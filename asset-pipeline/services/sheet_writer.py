"""
sheet_writer.py — Generate a multi-tab Excel .xlsx spreadsheet and
per-tab CSV files from structured asset data.

The tab structure is dynamic — it writes whatever tabs/headers/rows
are passed in from the parsed data, not a fixed schema.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def write_multi_tab_xlsx(
    tabs_data: dict[str, Any],
    output_path: str | Path,
) -> Path:
    """
    Create a multi-tab Excel spreadsheet at *output_path*.

    *tabs_data* should be a dict like:
    {
        "TabName": {
            "headers": ["col1", "col2", ...],
            "rows": [{"col1": "val", ...}, ...]
        },
        ...
    }

    Returns the resolved Path to the written file.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # Styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    data_alignment = Alignment(vertical="top", wrap_text=True)

    first_tab = True
    for tab_name, tab_info in tabs_data.items():
        headers = tab_info.get("headers", [])
        rows = tab_info.get("rows", [])

        if first_tab:
            ws = wb.active
            ws.title = tab_name[:31]  # Excel limits tab names to 31 chars
            first_tab = False
        else:
            ws = wb.create_sheet(title=tab_name[:31])

        # Write header row
        for col_idx, col_name in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data rows
        for row_idx, row_data in enumerate(rows, start=2):
            for col_idx, col_name in enumerate(headers, start=1):
                value = row_data.get(col_name, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = data_alignment
                cell.border = thin_border

        # Auto-size columns
        for col_idx, col_name in enumerate(headers, start=1):
            max_len = len(col_name)
            for row_idx in range(2, len(rows) + 2):
                val = str(ws.cell(row=row_idx, column=col_idx).value or "")
                max_len = max(max_len, min(len(val), 50))
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = max_len + 4

        # Freeze header row
        ws.freeze_panes = "A2"

    # Handle edge case: no tabs at all
    if first_tab:
        ws = wb.active
        ws.title = "Empty"
        ws.cell(row=1, column=1, value="No data")

    wb.save(output_path)
    return output_path.resolve()


def write_csv_zip(tabs_data: dict[str, Any]) -> bytes:
    """
    Write each tab as a separate CSV and return them bundled in a ZIP.

    Returns raw ZIP bytes suitable for a download response.
    """
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for tab_name, tab_info in tabs_data.items():
            headers = tab_info.get("headers", [])
            rows = tab_info.get("rows", [])

            csv_buf = io.StringIO()
            writer = csv.DictWriter(csv_buf, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow(row)

            safe_name = tab_name.replace("/", "-").replace("\\", "-")
            zf.writestr(f"{safe_name}.csv", csv_buf.getvalue())

    return buf.getvalue()
